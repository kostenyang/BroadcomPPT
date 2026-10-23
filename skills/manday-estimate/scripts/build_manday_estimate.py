#!/usr/bin/env python3
"""Build a man-day (人天) effort-estimate workbook from a JSON spec.

Two-level structure: per-task man-day estimates, a subtotal per phase, and a
grand total ("Total Estimated man/day"). Man-days only — no rates/cost.
Whole-day integers only (no decimals).

Usage:
    python build_manday_estimate.py spec.json -o output.xlsx [--lang zh|en|bi]

Single-sheet spec:
{
  "title": "...", "sheet_name": "WBS估算",
  "phases": [ {"wbs":"1","name":"前置","tasks":[
      {"wbs":"1.1","task":"...","manday":2,"notes":""} ]} ]
}

Multi-sheet spec (one tab per entry in "sheets"):
{
  "sheets": [
    {"sheet_name":"VCF→VCF", "title":"...", "phases":[ ... ]},
    {"sheet_name":"vSphere8→VCF", "title":"...", "phases":[ ... ]}
  ]
}
- manday is a whole number of days (integers only, no decimals). notes optional.
"""
import argparse, json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.properties import PageSetupProperties

# ---- Theme (white + blue / 藍綠紫 accents, matching the CTBC template) ----
CJK   = "Microsoft JhengHei"
NUM   = "Arial"
C_PHASE_BG = "1F3864"
C_PHASE_FG = "FFFFFF"
C_SUB_BG   = "D6E4F0"
C_HEAD_BG  = "2F5597"
C_TOTAL_BG = "FFFF00"
C_TOTAL_FG = "C00000"
C_INPUT    = "0000FF"
C_INPUT_BG = "FFF2CC"

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MD = "0"   # man-day number format — whole days only, no decimals

LABELS = {
    "zh": dict(wbs="WBS 編號", task="工作項目 / 任務", md="人天", notes="備註",
               subtotal="小計", total="Total Estimated man/day", unit="人天",
               tsuffix="— 人天估算 (Man-Day Estimate)"),
    "en": dict(wbs="WBS", task="Work Item / Task", md="Man-Days", notes="Notes",
               subtotal="Subtotal", total="Total Estimated man/day", unit="man-days",
               tsuffix="— Man-Day Estimate"),
}


def L(lang):
    if lang == "en":
        return LABELS["en"]
    if lang == "bi":
        z, e = LABELS["zh"], LABELS["en"]
        return {k: (f"{z[k]} / {e[k]}" if k not in ("total", "unit") else z[k]) for k in z}
    return LABELS["zh"]


def cjk_font(**kw):
    kw.setdefault("name", CJK)
    return Font(**kw)


def build_sheet(ws, sheet_spec, lab):
    """Render one man-day estimate onto worksheet ws."""
    title = sheet_spec.get("title", "專案")
    ws["A1"] = f'{title} {lab["tsuffix"]}'
    ws["A1"].font = cjk_font(bold=True, size=16, color=C_PHASE_BG)
    ws.merge_cells("A1:D1")
    ws.row_dimensions[1].height = 26

    HR = 3
    for j, h in enumerate([lab["wbs"], lab["task"], lab["md"], lab["notes"]], 1):
        c = ws.cell(HR, j, h)
        c.font = cjk_font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color=C_HEAD_BG)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[HR].height = 22

    subtotal_cells = []
    r = HR + 1
    for ph in sheet_spec.get("phases", []):
        label = f'{ph.get("wbs","")} {ph.get("name","")}'.strip()
        hc = ws.cell(r, 1, label)
        hc.font = cjk_font(bold=True, size=13, color=C_PHASE_FG)
        hc.alignment = Alignment(horizontal="left", vertical="center")
        for j in range(1, 5):
            cell = ws.cell(r, j)
            cell.fill = PatternFill("solid", start_color=C_PHASE_BG)
            cell.border = BORDER
        ws.row_dimensions[r].height = 24
        r += 1

        first_task = r
        for t in ph.get("tasks", []):
            a = ws.cell(r, 1, t.get("wbs", "")); a.font = Font(name=NUM, size=11)
            a.alignment = Alignment(horizontal="left", vertical="center"); a.border = BORDER
            b = ws.cell(r, 2, t.get("task", "")); b.font = cjk_font(size=11)
            b.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); b.border = BORDER
            md = t.get("manday", None)
            c = ws.cell(r, 3, md if md is not None else None)
            c.number_format = MD
            c.font = Font(name=NUM, size=11, color=C_INPUT)
            c.fill = PatternFill("solid", start_color=C_INPUT_BG)
            c.alignment = Alignment(horizontal="center", vertical="center"); c.border = BORDER
            d = ws.cell(r, 4, t.get("notes", "")); d.font = cjk_font(size=10)
            d.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); d.border = BORDER
            r += 1
        last_task = r - 1

        sc_label = ws.cell(r, 2, f'{ph.get("name","")} {lab["subtotal"]}'.strip())
        sc_label.font = cjk_font(bold=True, size=11)
        sc_label.alignment = Alignment(horizontal="right", vertical="center")
        sc_label.fill = PatternFill("solid", start_color=C_SUB_BG)
        ws.cell(r, 1).fill = PatternFill("solid", start_color=C_SUB_BG)
        sc = ws.cell(r, 3, f"=SUM(C{first_task}:C{last_task})" if last_task >= first_task else 0)
        sc.number_format = MD
        sc.font = Font(name=NUM, bold=True, size=11)
        sc.fill = PatternFill("solid", start_color=C_SUB_BG)
        sc.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 4).fill = PatternFill("solid", start_color=C_SUB_BG)
        for j in range(1, 5):
            ws.cell(r, j).border = BORDER
        subtotal_cells.append(f"C{r}")
        r += 1

    tlabel = ws.cell(r, 2, lab["total"])
    tlabel.font = Font(name=NUM, bold=True, size=12)
    tlabel.fill = PatternFill("solid", start_color=C_TOTAL_BG)
    tlabel.alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(r, 1).fill = PatternFill("solid", start_color=C_TOTAL_BG)
    tot = ws.cell(r, 3, ("=" + "+".join(subtotal_cells)) if subtotal_cells else 0)
    tot.number_format = MD
    tot.font = Font(name=NUM, bold=True, size=13, color=C_TOTAL_FG)
    tot.fill = PatternFill("solid", start_color=C_TOTAL_BG)
    tot.alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(r, 4).fill = PatternFill("solid", start_color=C_TOTAL_BG)
    for j in range(1, 5):
        ws.cell(r, j).border = BORDER
    ws.row_dimensions[r].height = 22

    for col, w in dict(A=12, B=72, C=12, D=40).items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"A{HR+1}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"{HR}:{HR}"


def build(spec, out_path, lang="zh"):
    lab = L(lang)
    wb = Workbook()
    # normalise to a list of sheet specs (backward compatible with single-sheet)
    if "sheets" in spec:
        sheets = spec["sheets"]
    else:
        sheets = [{"sheet_name": spec.get("sheet_name", "WBS估算"),
                   "title": spec.get("title", "專案"),
                   "phases": spec.get("phases", [])}]
    for i, sh in enumerate(sheets):
        name = (sh.get("sheet_name") or f"WBS{i+1}")[:31]
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        build_sheet(ws, sh, lab)
    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("-o", "--out", default="manday_estimate.xlsx")
    ap.add_argument("--lang", default="zh", choices=["zh", "en", "bi"])
    a = ap.parse_args()
    spec = json.load(open(a.spec, encoding="utf-8"))
    print("saved", build(spec, a.out, lang=a.lang))


if __name__ == "__main__":
    main()
